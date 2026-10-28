"""استئناف بعد انقطاع مفاجئ + تصدير Word: python -m unittest discover -s tests"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
os.environ.setdefault("PDF2MD_DATA_DIR", tempfile.mkdtemp(prefix="pdf2md-test-data-"))

from app import checkpoint as cp  # noqa: E402
from app import pipeline  # noqa: E402
from app.config import Settings  # noqa: E402
from app.docx_out import markdown_to_docx_bytes  # noqa: E402
from app.hashing import HashCache  # noqa: E402
from app.pipeline import BatchContext, Job, run_batch  # noqa: E402
from make_sample_pdf import build  # noqa: E402


class CrashResumeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.pdf = build(self.tmp / "in" / "crash.pdf")
        self.out = self.tmp / "out"
        self.cache = HashCache(self.tmp / "cache.json")
        self.job = Job(self.pdf, self.cache.get(self.pdf))

    def _ctx(self, stop=None):
        return BatchContext(self.out, Settings(), stop or threading.Event(), lambda *a, **k: None, self.cache,
                            beside_source=True)

    def test_hard_crash_leaves_snapshot_and_progress(self):
        """انقطاع بلا finally (كهرباء/قتل العملية): الأجزاء محفوظة، ولقطة Markdown موجودة، والتقدّم معروف."""
        real = pipeline._process_page
        calls = []

        def crash_on_third(*a, **k):
            if len(calls) == 2:
                raise SystemExit("power off")
            calls.append(1)
            real(*a, **k)

        with mock.patch.object(pipeline, "SNAPSHOT_EVERY", 1), \
                mock.patch.object(pipeline, "_process_page", crash_on_third), \
                mock.patch.object(pipeline, "finalize_file", side_effect=SystemExit("power off")):
            with self.assertRaises(SystemExit):
                run_batch([self.job], self._ctx())
        md = self.pdf.parent / "crash.md"
        self.assertTrue(md.is_file(), "لا لقطة Markdown بعد الانقطاع")
        self.assertIn("<!-- page:1 -->", md.read_text(encoding="utf-8"))
        done, total = cp.progress(self.out / ".state", self.job.file_hash) or (0, 0)
        self.assertEqual((done, total), (1, 5))   # الصفحة 1 نصية، والثانية مصوّرة بلا محرّك (غير نهائية)

        # بعد إعادة التشغيل: تكتمل الصفحات المتبقية فقط
        run_batch([self.job], self._ctx())
        st = cp.FileState(self.out / ".state", self.job.file_hash, "crash.pdf", 5)
        self.assertEqual(sorted(st.data["pages"]), ["1", "2", "3", "4", "5"])

    def test_split_parts_survive_crash_instead_of_growing_file(self):
        real = pipeline._process_page
        calls = []

        def crash_on_third(*a, **k):
            if len(calls) == 2:
                raise SystemExit("power off")
            calls.append(1)
            real(*a, **k)

        ctx = self._ctx()
        ctx.settings.split_parts, ctx.settings.split_parts_pages = True, 1
        with mock.patch.object(pipeline, "SNAPSHOT_EVERY", 1),                 mock.patch.object(pipeline, "_process_page", crash_on_third),                 mock.patch.object(pipeline, "finalize_file", side_effect=SystemExit("power off")):
            with self.assertRaises(SystemExit):
                run_batch([self.job], ctx)
        folder = self.pdf.parent
        self.assertTrue((folder / "crash.part01.md").is_file())
        self.assertFalse((folder / "crash.md").exists())   # لا ملف كامل يكبر أثناء المعالجة

    def test_progress_none_for_unknown_and_corrupt(self):
        self.assertIsNone(cp.progress(self.out / ".state", "nohash"))
        (self.out / ".state").mkdir(parents=True)
        (self.out / ".state" / "bad.json").write_text("{not json", encoding="utf-8")
        self.assertIsNone(cp.progress(self.out / ".state", "bad"))


class DocxTests(unittest.TestCase):
    MD = ("# عقد\n\n<!-- page:1 -->\n**[صفحة 1]**\n\nنص **عريض** عربي 250,000 دينار.\n\n"
          "| الاسم | المبلغ |\n| --- | --- |\n| أحمد | 75,600 |\n\n- بند\n\n## English\nPlain *text* here.\n")

    def _xml(self, md):
        with zipfile.ZipFile(__import__("io").BytesIO(markdown_to_docx_bytes(md))) as z:
            self.assertIsNone(z.testzip())
            return z.read("word/document.xml").decode("utf-8")

    def test_arabic_is_rtl_everywhere(self):
        x = self._xml(self.MD)
        self.assertIn("<w:bidi/>", x)            # الفقرات والقسم
        self.assertIn("<w:bidiVisual/>", x)      # الجدول: العمود الأول في اليمين
        self.assertIn("<w:rtl/>", x)
        self.assertIn('<w:jc w:val="both"/>', x)
        self.assertNotIn("<!--", x)

    def test_english_stays_ltr(self):
        x = self._xml("# Title\n\nPlain english text only.\n\n| a | b |\n| - | - |\n| 1 | 2 |\n")
        self.assertNotIn("<w:bidi/>", x)
        self.assertNotIn("<w:bidiVisual/>", x)

    def test_mixed_paragraph_follows_majority(self):
        x = self._xml("English sentence that is long enough with one كلمة only.\n")
        self.assertNotIn("<w:bidi/>", x)

    def test_opens_with_python_docx_if_available(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx غير مثبّت")
        import io
        d = docx.Document(io.BytesIO(markdown_to_docx_bytes(self.MD)))
        self.assertEqual(len(d.tables), 1)
        self.assertEqual([c.text for c in d.tables[0].rows[1].cells], ["أحمد", "75,600"])
        self.assertTrue(any("عريض" in p.text for p in d.paragraphs))

    def test_pipeline_exports_docx_when_enabled(self):
        tmp = Path(tempfile.mkdtemp())
        pdf = build(tmp / "in" / "w.pdf")
        cache = HashCache(tmp / "c.json")
        s = Settings()
        s.export_docx = True
        ctx = BatchContext(tmp / "out", s, threading.Event(), lambda *a, **k: None, cache, beside_source=True)
        run_batch([Job(pdf, cache.get(pdf))], ctx)
        self.assertTrue((pdf.parent / "w.docx").is_file())
        self.assertTrue(re.search(rb"<w:bidi/>", zipfile.ZipFile(pdf.parent / "w.docx").read("word/document.xml")))


if __name__ == "__main__":
    unittest.main()


class SplitParts(unittest.TestCase):
    def _state(self, n=7, done=(1, 2, 3, 4, 5)):
        import tempfile
        from app import checkpoint as cp
        d = Path(tempfile.mkdtemp())
        st = cp.FileState(d / "state", "h" * 8, "كتاب.pdf", n)
        for p in done:
            st.save_page(p, cp.DONE, "text", "نص صفحة %d" % p)
        return d, st

    def test_parts_written_only_for_complete_blocks_and_removed(self):
        from app import markdown_out as mo
        d, st = self._state()
        out = d / "out"
        out.mkdir()
        w = mo.write_ready_parts(st, "كتاب.pdf", out, 2)
        self.assertEqual([p.name for p in w], ["كتاب.part01.md", "كتاب.part02.md"])   # ص1-2 و3-4؛ ص5-6 ناقصة
        text = (out / "كتاب.part02.md").read_text(encoding="utf-8")
        self.assertIn("(ص3–4)", text)
        self.assertIn("نص صفحة 3", text)
        self.assertNotIn("نص صفحة 5", text)
        self.assertEqual(mo.write_ready_parts(st, "كتاب.pdf", out, 2), [])   # لا يعاد كتابة الجزء المكتوب
        self.assertEqual(mo.remove_parts(out, "كتاب.pdf"), 2)
        self.assertEqual(list(out.glob("*.md")), [])
